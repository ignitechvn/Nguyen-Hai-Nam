"""Convert the website layout spreadsheet into src/content.json.

Usage: python3 tools/xlsx_to_json.py [content/Website_Layout_Hai_Nam.xlsx] [src/content.json]
Requires: pip install openpyxl

Photos are listed in content/photos.json, keyed by the start of a record title.
"""
import json
import re
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else 'content/Website_Layout_Hai_Nam.xlsx'
DEST = sys.argv[2] if len(sys.argv) > 2 else 'src/content.json'
PHOTOS = 'content/photos.json'

# Columns: A No | B Categories | C Breakdown | D Content [VIET] | E Content [ENG] | F Image/Video | G Notes
VI_CHARS = re.compile(r'[ăâđêôơưàảãáạằẳẵắặầẩẫấậèẻẽéẹềểễếệìỉĩíịòỏõóọồổỗốộờởỡớợùủũúụừửữứựỳỷỹýỵ]', re.I)
SMALL_WORDS = {'a', 'an', 'and', 'the', 'to', 'of', 'in', 'on', 'for', 'with'}
ACRONYMS = {'VEX', 'AI', 'STEM', 'IOT', 'HKU', 'NTU'}
CATEGORY_NAMES = {'Home': 'About Me'}  # page name in the sheet -> menu name on the site


def is_vietnamese(text):
    return bool(VI_CHARS.search(text or ''))


def title_case(text):
    words = text.strip().split()
    out = []
    for i, w in enumerate(words):
        lw = w.lower()
        if w.upper() in ACRONYMS:
            out.append(w.upper())
        elif i > 0 and lw in SMALL_WORDS:
            out.append(lw)
        else:
            out.append(lw[:1].upper() + lw[1:])
    return ' '.join(out)


def clean_text(text):
    """Split content into paragraphs, drop designer placeholders like [Insert ảnh ...]."""
    if not text:
        return []
    text = text.strip().strip('"“”').strip()
    paragraphs = []
    for line in text.split('\n'):
        line = line.strip()
        if not line or re.fullmatch(r'\[[^\]]*\]', line):
            continue
        line = re.sub(r'\s*\[Insert[^\]]*\]\s*', ' ', line).strip()
        if line.endswith('| Certificate'):
            continue  # certificate list, shown as captioned images instead
        if line:
            paragraphs.append(line)
    return paragraphs


def link_label(note):
    note = (note or '').lower()
    if 'nghiên cứu' in note:
        return 'research'
    if 'fanpage' in note:
        return 'fanpage'
    if 'công ty' in note:
        return 'company'
    return 'platform'


def hyperlink(cell):
    return cell.hyperlink.target if cell.hyperlink else ''


def main():
    ws = openpyxl.load_workbook(SRC).active

    profile_lines = (ws['A2'].value or '').split('\n')
    contact = dict(re.findall(r'(Mail|Number|Domain):\s*([^|]+)', profile_lines[1] if len(profile_lines) > 1 else ''))
    title_parts = (ws['A1'].value or '').split('|')
    profile = {
        'name': title_parts[-1].split('-')[0].strip(),
        'major': profile_lines[0].replace('Major:', '').strip(),
        'email': contact.get('Mail', '').strip(),
        'phone': contact.get('Number', '').strip(),
        'domain': contact.get('Domain', '').strip(),
    }

    with open(PHOTOS, encoding='utf-8') as fp:
        photos = json.load(fp)

    # Sub-group headings span columns C:G (e.g. "OTHER ROBOTICS EXPERIENCES")
    group_rows = {r.min_row for r in ws.merged_cells.ranges if r.min_col == 3 and r.max_col == 7}

    records = []
    category = ''
    group = ''
    pending_title = ''

    for row in range(5, ws.max_row + 1):
        a, b, c, d, e, f, g = (ws.cell(row, col) for col in range(1, 8))

        if isinstance(a.value, str) and 'Page:' in a.value:
            category = title_case(a.value.split(':', 1)[1])
            category = CATEGORY_NAMES.get(category, category)
            group = ''
            pending_title = ''
            continue

        if row in group_rows and c.value:
            group = c.value.split('\n')[0].strip()
            group = title_case(group) if group.isupper() else group

        if not c.value:
            continue

        if 'Không cần content' in str(d.value or ''):
            vi, en = [], []  # designer note, not page content
        else:
            vi = clean_text(d.value)
            en = clean_text(e.value)

        has_photos = any(c.value.startswith(key) for key in photos)
        if not vi and not en and not has_photos:
            # Section heading without content, used as title for the next row
            pending_title = c.value.split('\n')[0].strip()
            continue

        lines = [l.strip() for l in c.value.split('\n') if l.strip()]
        # A quoted title may continue on the next lines
        while len(lines) > 1 and lines[0].count('“') > lines[0].count('”'):
            lines[0:2] = [lines[0] + ' ' + lines[1]]
        heading, subtitle = lines[0], ' '.join(lines[1:])
        if is_vietnamese(heading):
            # Breakdown column holds an instruction, not a title
            heading = b.value or pending_title or heading
            subtitle = ''
        if heading.isupper():
            heading = title_case(heading)

        role, _, title = heading.partition('|')
        if not title:
            role, title = '', heading

        links = []
        if hyperlink(g):
            links.append({'type': link_label(g.value), 'url': hyperlink(g)})

        records.append({
            'category': category,
            'group': group,
            'title': title.strip(),
            'role': role.strip(),
            'subtitle': subtitle,
            'en': en or vi,
            'vi': vi or en,
            'links': links,
        })
        pending_title = ''

    for key, entry in photos.items():
        matches = [r for r in records if r['title'].startswith(key)]
        if not matches:
            print(f'warning: no record for photos "{key}"')
        for record in matches:
            record.update(entry)

    with open(DEST, 'w', encoding='utf-8') as fp:
        json.dump({'profile': profile, 'records': records}, fp, ensure_ascii=False, indent=2)
    print(f'{len(records)} records -> {DEST}')


if __name__ == '__main__':
    main()
